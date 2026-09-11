#!/usr/bin/env python3
"""convert a collaborator's one-place evidence spreadsheet into an
agent-dossier.v1 json file, quarantining personal details at capture.

input: the three-sheet workbook joseph watts produced for st martin's,
loburn (summary; evidence timeline; notes). output: a dossier whose
provenance says collaborator_import, whose claims carry the row's locator
and any verbatim fragment the row quoted, and whose quarantine block holds
every phone, email and honorific-led name the rows contained.

usage:
  uv run --with openpyxl python scripts/agent_research/import_watts_xlsx.py \
      /path/to/St_Martins_Loburn_PoW_Evidence.xlsx out.dossier.json [--redact]

--redact strips the quarantined values (hashes remain); use it for any copy
that leaves the run directory.
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import lib  # noqa: E402

PROMPT_VERSION = "collaborator-xlsx-import-v1"
SEED_SOURCE = "nz_places.geojson (overpass extract 2025-08-20)"

# jw's summary said glm 5.3 flash through openrouter; the sheet itself names no model
COLLABORATOR_MODEL_ID = "z-ai/glm-5.3-flash (collaborator-reported)"

SOURCE_TYPES = [
    ("openstreetmap", "osm"),
    ("obituary", "newspaper"),
    ("press", "newspaper"),
    ("flickr", "photograph_caption"),
    ("canterbury stories", "archive_collection"),
    ("directory", "directory"),
    ("e-life", "diocesan_publication"),
    ("parish", "church_website"),
    ("anglican life rangiora", "church_website"),
]

WEIGHTS = {
    "osm": "user_contributed",
    "newspaper": "contemporary_report",
    "photograph_caption": "user_contributed",
    "archive_collection": "secondary",
    "directory": "primary_institutional",
    "diocesan_publication": "primary_institutional",
    "church_website": "primary_institutional",
}

# keyword rules from the row text to claim types; one row may yield several
CLAIM_RULES = [
    ("building_date", re.compile(r"\b(built|constructed|construction)\b.*\b(1[6-9]\d\d|20\d\d)\b", re.I)),
    ("renovation", re.compile(r"\brenovat", re.I)),
    ("land_or_consecration", re.compile(r"\b(consecrated|acquired land)\b", re.I)),
    ("worship_active_asof", re.compile(r"\b(funeral service|functioning as|services were held)\b", re.I)),
    ("service_pattern", re.compile(r"\b(service_times|services were held|10:00|10 am)\b", re.I)),
    ("sale_or_disposal", re.compile(r"\bapproved for sale\b", re.I)),
    ("worship_ended", re.compile(r"\bfinal services?\b", re.I)),
    ("closure_event", re.compile(r"\bclosed\b", re.I)),
    ("organisation_link", re.compile(r"\b(parish|diocese)\b.*\b(identifies|relationship)\b", re.I)),
    ("address", re.compile(r"\b(street address|\d+ [A-Z][a-z]+ (Road|Street))\b")),
]

_MONTHS = {m.lower(): i for i, m in enumerate(
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], start=1)}
_DMY = re.compile(r"\b(\d{1,2}) (Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]* (\d{4})\b")
_YEAR = re.compile(r"\b(1[6-9]\d\d|20\d\d)\b")
_VERSION = re.compile(r"Version (\d+)")
_CHANGESET = re.compile(r"changeset (\d+)")
_CURLY_QUOTES = re.compile(r"[“\"]([^”\"]{6,400})[”\"]")


def parse_date(text: str) -> tuple[str | None, str, str]:
    """(iso date, precision, basis) from a source-date cell"""
    if not text:
        return None, "unknown", "not_stated"
    lowered = text.lower()
    m = _DMY.search(text)
    if m:
        iso = f"{int(m.group(3)):04d}-{_MONTHS[m.group(2).lower()[:3]]:02d}-{int(m.group(1)):02d}"
        basis = "metadata" if "metadata" in lowered or "uploaded" in lowered else "printed"
        return iso, "day", basis
    y = _YEAR.search(text)
    if y:
        precision = "approximate" if "approx" in lowered or "late" in lowered or "early" in lowered else "year"
        return y.group(1), precision, "inferred" if precision == "approximate" else "printed"
    return None, "unknown", "not_stated"


DOMAIN_TYPES = [
    ("openstreetmap.org", "osm"),
    ("paperspast.natlib.govt.nz", "newspaper"),
    ("deaths.press.co.nz", "newspaper"),
    ("flickr.com", "photograph_caption"),
    ("canterburystories.nz", "archive_collection"),
    ("anglicanlife.org.nz/church/", "directory"),
    ("anglicanlife.org.nz/anglican-e-life", "diocesan_publication"),
    ("anglicanliferangiora.church", "church_website"),
    ("heritage.org.nz", "heritage_register"),
    ("wikidata.org", "wikidata"),
    ("wikipedia.org", "wikipedia"),
    ("charities.govt.nz", "charity_register"),
]


def source_type_for(source_name: str, locator: str = "") -> str:
    # the locator's domain is a steadier guide than the collaborator's label
    lowered_url = (locator or "").lower()
    for needle, stype in DOMAIN_TYPES:
        if needle in lowered_url:
            return stype
    lowered = (source_name or "").lower()
    for needle, stype in SOURCE_TYPES:
        if needle in lowered:
            return stype
    return "other"


# where a row carries several years, the year that belongs to the claim type
TYPE_YEAR = {
    "building_date": re.compile(r"\b(?:built|constructed|construction)\b[^.;]*?\b(1[6-9]\d\d|20\d\d)\b", re.I),
    "renovation": re.compile(r"\brenovat\w*\b[^.;]*?\b(1[6-9]\d\d|20\d\d)\b", re.I),
    "land_or_consecration": re.compile(r"\b(?:acquired land|consecrated)\b[^.;]*?\b(1[6-9]\d\d|20\d\d)\b", re.I),
}


def year_for(ctype: str, text: str) -> str | None:
    rule = TYPE_YEAR.get(ctype)
    if rule:
        m = rule.search(text)
        if m:
            return m.group(1)
    years = _YEAR.findall(text)
    return years[0] if years else None


def quoted_fragments(*cells: str) -> str:
    fragments = []
    for cell in cells:
        for m in _CURLY_QUOTES.finditer(cell or ""):
            fragments.append(m.group(1).strip())
    return " … ".join(fragments)[:600]


def read_sheets(path: Path) -> tuple[dict, list[dict], dict]:
    import openpyxl  # optional dependency, only for the import

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    summary = {}
    for row in wb["Summary"].iter_rows(values_only=True):
        if row and row[0] and row[1] is not None:
            summary[str(row[0]).strip()] = str(row[1]).strip()
    rows = []
    sheet = wb["Evidence Timeline"]
    header = None
    for row in sheet.iter_rows(values_only=True):
        if header is None:
            header = [str(h).strip() for h in row]
            continue
        if not any(row):
            continue
        rows.append({header[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
    notes = {}
    for row in wb["Notes"].iter_rows(values_only=True):
        if row and row[0] and len(row) > 1 and row[1]:
            notes[str(row[0]).strip()] = str(row[1]).strip()
    return summary, rows, notes


def build_dossier(summary: dict, rows: list[dict], notes: dict, place_ref: str,
                  seed_lat: float, seed_lon: float, produced_by: str, source_file: str,
                  extra_names: list[str] | None = None) -> dict:
    retrieved_default = None
    assessment_date = summary.get("Assessment date", "")
    iso, _, _ = parse_date(assessment_date)
    asof = iso or datetime.utcnow().strftime("%Y-%m-%d")
    reader = {
        "reader_id": f"collaborator:{produced_by}",
        "backend": "collaborator_import",
        "model_id": COLLABORATOR_MODEL_ID,
        "prompt_version": PROMPT_VERSION,
    }
    claims = []
    chain = []
    counter = 0

    def claim_id():
        nonlocal counter
        counter += 1
        return f"{place_ref}:c{counter:02d}"

    for row in rows:
        record_type = row.get("Record type", "")
        version_cell = row.get("Version / source date", "")
        state = row.get("Complete state or information supplied", "")
        change = row.get("Change / additional information supplied", "")
        source_name = row.get("Source name", "")
        locator = row.get("Source URL", "")
        retrieved, _, _ = parse_date(row.get("Retrieved", ""))
        retrieved = retrieved or retrieved_default or asof
        stype = source_type_for(source_name, locator)
        source_date, precision, basis = parse_date(version_cell)
        source = {
            "locator": locator,
            "source_name": source_name,
            "source_type": stype,
            "source_date": source_date,
            "source_date_basis": basis,
            "retrieved_at": retrieved,
        }
        if record_type.lower().startswith("osm edit"):
            version = int(_VERSION.search(version_cell).group(1)) if _VERSION.search(version_cell) else None
            changeset = int(_CHANGESET.search(version_cell).group(1)) if _CHANGESET.search(version_cell) else None
            chain.append({
                "version": version,
                "changeset": changeset,
                "timestamp": source_date,
                "timestamp_basis": "approximate",
                "tags_summary": state,
                "change_note": change,
                "locator": locator,
            })
            claims.append({
                "claim_id": claim_id(),
                "claim_type": "osm_object_version",
                "value": f"version {version}, changeset {changeset}",
                "value_structured": {"version": version, "changeset": changeset},
                "date_start": source_date,
                "date_end": None,
                "date_precision": precision,
                "source": source,
                "quoted_support": quoted_fragments(change),
                "evidential_weight": "user_contributed",
                "reader": reader,
                "confidence": "medium",
                "note": change,
            })
            continue
        text = f"{state} {change}"
        matched = [ctype for ctype, rule in CLAIM_RULES if rule.search(text)]
        if not matched:
            matched = ["other"]
        weight = WEIGHTS.get(stype, "secondary")
        if "moderate evidential weight" in change.lower() or "requiring corroboration" in change.lower():
            weight = "user_contributed"
        for ctype in matched:
            date_start = source_date if ctype in ("worship_active_asof", "sale_or_disposal", "worship_ended", "closure_event") else year_for(ctype, state)
            claims.append({
                "claim_id": claim_id(),
                "claim_type": ctype,
                "value": state,
                "date_start": date_start,
                "date_end": None,
                "date_precision": precision if date_start == source_date else ("year" if date_start else "unknown"),
                "source": source,
                "quoted_support": quoted_fragments(state, change),
                "evidential_weight": weight,
                "reader": reader,
                "confidence": "high" if weight == "primary_institutional" else "medium",
                "note": change,
            })

    status_text = summary.get("Current status assessment", "").lower()
    if "likely inactive" in status_text or "closed" in status_text:
        status = "likely_inactive"
    elif "inactive" in status_text:
        status = "inactive"
    elif "active" in status_text:
        status = "likely_active"
    else:
        status = "unknown"
    supporting = [c["claim_id"] for c in claims if c["claim_type"] in ("closure_event", "worship_ended", "sale_or_disposal")]
    osm_status = summary.get("OSM status", "")
    dossier = {
        "schema_version": lib.SCHEMA_VERSION,
        "dossier_id": f"{place_ref}:{reader['reader_id']}:{asof}",
        "place": {
            "place_ref": place_ref,
            "name": summary.get("Name", ""),
            "country_code": "NZ",
            "seed_latitude": seed_lat,
            "seed_longitude": seed_lon,
            "seed_source": SEED_SOURCE,
        },
        "candidate_location": {
            "latitude": None,
            "longitude": None,
            "basis": "not_assessed",
            "basis_note": "the spreadsheet gives an address and the osm object but no coordinate of its own",
            "uncertainty_radius_m": None,
            "address": summary.get("Address"),
        },
        "claims": claims,
        "status_assessment": {
            "current_status": status,
            "basis": summary.get("Status evidence", ""),
            "asof_date": asof,
            "supporting_claim_ids": supporting,
            "osm_stale": osm_status.lower().startswith("appears stale") or None,
            "osm_stale_basis": osm_status,
        },
        "osm_version_chain": chain,
        "personal_details_quarantine": {"redacted": False, "item_count": 0, "items": []},
        "provenance": {
            "producer": "collaborator_import",
            "ai_generated": True,
            "produced_by": produced_by,
            "imported_from": f"{source_file} (Summary; Evidence Timeline; Notes)",
            "lane": "agent_assisted",
        },
        "run_manifest": {
            "run_id": f"import-{asof}",
            "backend": "collaborator_import",
            "model_id_requested": COLLABORATOR_MODEL_ID,
            "model_id_reported": None,
            "prompt_version": PROMPT_VERSION,
            "started_at": asof,
            "ended_at": asof,
            "duration_s": None,
            "usage": {},
            "cost_usd_reported": None,
            "cost_basis": "collaborator_reported",
            "idempotency_key": lib.idempotency_key(place_ref, PROMPT_VERSION, COLLABORATOR_MODEL_ID, SEED_SOURCE),
            "tool_permissions": [],
            "exit_status": "completed",
            "notes": " | ".join(f"{k}: {v}" for k, v in notes.items()),
        },
    }
    lib.quarantine_dossier(dossier, extra_names=extra_names)
    return dossier


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("out", type=Path)
    parser.add_argument("--place-ref", default="osm:way/643590665")
    parser.add_argument("--seed-lat", type=float, default=-43.24923163636364)
    parser.add_argument("--seed-lon", type=float, default=172.52855542727272)
    parser.add_argument("--produced-by", default="JW")
    parser.add_argument("--quarantine-name", action="append", default=[],
                        help="a person's name to quarantine even when no honorific precedes it (repeatable)")
    parser.add_argument("--redact", action="store_true", help="strip quarantined values before writing")
    args = parser.parse_args(argv)

    summary, rows, notes = read_sheets(args.xlsx)
    dossier = build_dossier(summary, rows, notes, args.place_ref, args.seed_lat, args.seed_lon,
                            args.produced_by, args.xlsx.name, extra_names=args.quarantine_name)
    if args.redact:
        lib.redact_quarantine(dossier)
    errors = lib.validate_dossier(dossier)
    if errors:
        for error in errors:
            print("schema:", error, file=sys.stderr)
        return 1
    lib.write_json(args.out, dossier)
    print(f"wrote {args.out}: {len(dossier['claims'])} claims, {len(dossier['osm_version_chain'])} osm versions, "
          f"{dossier['personal_details_quarantine']['item_count']} personal details quarantined"
          f"{' (redacted)' if args.redact else ''}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
